"""Builds data.js for the EMS dashboard from the KWH report workbook.
Usage:  python build_data.py [path/to/KWH_REPORT.xlsx]
Needs:  pip install openpyxl pandas numpy
Note: the workbook must have been saved by Excel so formula results are stored."""
import sys, re, json, datetime, collections
import numpy as np, pandas as pd
from openpyxl import load_workbook

SRC = sys.argv[1] if len(sys.argv) > 1 else 'KWH_REPORT.xlsx'
PANELS = ['Busway 1','Busway 2','MCC-1','MCC-2','MCC-3','MCC-4','MCC-5','LPDB-1','LPDB-2','LPDB Office']

def norm_panel(label):
    s = re.sub(r'\s+', ' ', str(label)).upper()
    if 'BUSWAY 1' in s: return 'Busway 1'
    if 'BUSWAY 2' in s: return 'Busway 2'
    m = re.search(r'MCC[- ]?(\d)', s)
    if m: return f'MCC-{m.group(1)}'
    m = re.search(r'LPDB[- ]?(\d)', s)
    if m: return f'LPDB-{m.group(1)}'
    if 'LPDB' in s and 'OFFICE' in s: return 'LPDB Office'
    return None

def norm_tail(sub, grp):
    s = (str(sub) + ' ' + str(grp or '')).lower()
    if 'export' in s: return 'solar_export'
    if 'loss' in s or 'dif b' in s or 'diff b' in s: return 'losses'
    if 'consumed by fb' in s: return 'fb_kwh'
    if 'solar' in s: return 'solar2' if '#2' in s else 'solar'
    if re.search(r'\blv\b', s): return 'main_lv'
    if re.search(r'\bmv\b', s): return 'main_mv'
    if 'main incoming' in s: return 'main_lv'
    return None

def read_workbook(path):
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        op = next((i for i, r in enumerate(rows[:20]) if any(isinstance(c, str) and c.strip().lower() == 'opening' for c in r)), None)
        if op is None: continue
        hdr = rows[op]; lab = rows[op-1] if op > 0 else (); sub = rows[op+1] if op+1 < len(rows) else ()
        cols = {}
        for c, v in enumerate(hdr):
            if isinstance(v, str) and v.strip().lower() == 'opening':
                name = norm_panel(lab[c] if c < len(lab) else None)
                if name: cols.setdefault(name, c+2)
        for c, v in enumerate(sub):
            if isinstance(v, str) and v.strip():
                g = next((hdr[k] for k in range(c, max(c-3, -1), -1) if k < len(hdr) and isinstance(hdr[k], str) and hdr[k].strip()), None)
                t = norm_tail(v, g)
                if t and v.strip().lower() not in ('opening', 'closing') and t not in cols: cols[t] = c
        dcol = next((c for c in range(3) if any(isinstance(r[c] if c < len(r) else None, datetime.datetime) for r in rows)), None)
        data = []
        for r in rows:
            d = r[dcol] if dcol is not None and dcol < len(r) else None
            if not isinstance(d, datetime.datetime): continue
            rec = {'date': d.date()}
            for k, c in cols.items():
                v = r[c] if c < len(r) else None
                rec[k] = float(v) if isinstance(v, (int, float)) else None
            data.append(rec)
        if data:  # keep the sheet's own month
            ym = collections.Counter((x['date'].year, x['date'].month) for x in data).most_common(1)[0][0]
            out += [x for x in data if (x['date'].year, x['date'].month) == ym]
    return pd.DataFrame(out)

df = read_workbook(SRC)
df['date'] = pd.to_datetime(df['date']); df = df.sort_values('date').drop_duplicates('date').reset_index(drop=True)
for c in PANELS + ['main_lv', 'solar', 'solar2', 'fb_kwh']:
    if c not in df: df[c] = np.nan
log = []

# drop days after the last recorded reading (empty rows give zero / huge negative formulas)
psum = df[PANELS].sum(axis=1, min_count=1)
df = df[df.date <= df.loc[psum > 0, 'date'].max()].copy()

# months where only a few panels were recorded are incomplete
ym = df.date.dt.to_period('M')
for p, g in df.groupby(ym):
    if g[PANELS].notna().any().sum() < 5:
        df.loc[g.index, PANELS + ['fb_kwh', 'main_lv']] = np.nan
        log.append([str(p), 'Month incomplete in source (too few panels recorded) – excluded'])

# obvious reading errors: > 10x local median (Busways swap load, so they are skipped)
for c in [x for x in PANELS if not x.startswith('Busway')] + ['main_lv']:
    s = df[c]; med = s.where(s > 0).rolling(31, center=True, min_periods=5).median()
    for i in df.index[(s > 10*med) & (med > 50)]:
        log.append([df.at[i, 'date'].strftime('%Y-%m-%d'), f"{'Main incoming (LV)' if c=='main_lv' else c}: {s[i]:,.0f} kWh looks like a reading error (typical ≈ {med[i]:,.0f}) – replaced with local median"])
        df.at[i, c] = round(med[i], 1)

panel_sum = df[PANELS].sum(axis=1, min_count=1)
df['total'] = df['fb_kwh'].where(df['fb_kwh'].notna(), panel_sum)
df.loc[panel_sum.isna() & df.fb_kwh.isna(), 'total'] = np.nan
df['solar_t'] = df[['solar', 'solar2']].sum(axis=1, min_count=1)
df.loc[df.main_lv.abs() > 3*df.total.fillna(1e9), 'main_lv'] = np.nan

r = lambda x: None if pd.isna(x) else round(float(x), 1)
rows = [[d.strftime('%Y-%m-%d'), r(t)] + [r(x) for x in v] + [r(g), r(s)]
        for d, t, v, g, s in zip(df.date, df.total, df[PANELS].values, df.main_lv, df.solar_t)]
out = {'panels': PANELS, 'rows': rows, 'notes': log, 'last': df.date.max().strftime('%Y-%m-%d')}
open('data.js', 'w', encoding='utf-8').write('window.EMS_DATA=' + json.dumps(out, separators=(',', ':'), ensure_ascii=False) + ';\n')
print(f"data.js written: {len(rows)} days, {df.date.min().date()} to {df.date.max().date()}")
